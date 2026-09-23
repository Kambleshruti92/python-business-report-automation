from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_report(employee_data, department_summary):
    """
    Generate a professional Excel report containing
    employee performance and department summary.
    """

    output_file = "reports/employee_performance_report.xlsx"

    workbook = Workbook()

    # --------------------------------------------------
    # Common border style
    # --------------------------------------------------

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # ==================================================
    # 1. Employee Performance Sheet
    # ==================================================

    employee_sheet = workbook.active
    employee_sheet.title = "Employee Performance"

    # Report title
    employee_sheet["A1"] = "Employee Performance Report"

    employee_sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    employee_sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    employee_sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(employee_data.columns)
    )

    # Headers
    employee_headers = list(employee_data.columns)

    employee_sheet.append(employee_headers)

    # Header formatting
    for cell in employee_sheet[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Employee data
    for row in employee_data.itertuples(
        index=False,
        name=None
    ):
        employee_sheet.append(row)

    # Add borders to employee data
    for row in employee_sheet.iter_rows(
        min_row=3,
        max_row=employee_sheet.max_row,
        min_col=1,
        max_col=employee_sheet.max_column
    ):
        for cell in row:
            cell.border = thin_border

    # Format Completion Percentage column
    for row in range(3, employee_sheet.max_row + 1):
        employee_sheet.cell(
            row=row,
            column=7
        ).number_format = "0.00"

    # Freeze headers
    employee_sheet.freeze_panes = "A3"

    # Enable filtering
    employee_sheet.auto_filter.ref = (
        f"A2:"
        f"{get_column_letter(employee_sheet.max_column)}"
        f"{employee_sheet.max_row}"
    )

    # ==================================================
    # 2. Department Summary Sheet
    # ==================================================

    summary_sheet = workbook.create_sheet(
        "Department Summary"
    )

    # Report title
    summary_sheet["A1"] = "Department Performance Summary"

    summary_sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    summary_sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(department_summary.columns)
    )

    # Headers
    summary_headers = list(
        department_summary.columns
    )

    summary_sheet.append(summary_headers)

    # Header formatting
    for cell in summary_sheet[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Department data
    for row in department_summary.itertuples(
        index=False,
        name=None
    ):
        summary_sheet.append(row)

    # Add borders to department data
    for row in summary_sheet.iter_rows(
        min_row=3,
        max_row=summary_sheet.max_row,
        min_col=1,
        max_col=summary_sheet.max_column
    ):
        for cell in row:
            cell.border = thin_border

    # Format Completion Percentage column
    for row in range(3, summary_sheet.max_row + 1):
        summary_sheet.cell(
            row=row,
            column=6
        ).number_format = "0.00"

    # Freeze headers
    summary_sheet.freeze_panes = "A3"

    # Enable filtering
    summary_sheet.auto_filter.ref = (
        f"A2:"
        f"{get_column_letter(summary_sheet.max_column)}"
        f"{summary_sheet.max_row}"
    )

    # ==================================================
    # 3. Automatically Adjust Column Widths
    # ==================================================

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[
                column_letter
            ].width = max_length + 2

    # ==================================================
    # 4. Save Excel Report
    # ==================================================

    workbook.save(output_file)

    return output_file